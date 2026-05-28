from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from datetime import datetime
import pathlib
import os.path


class ExctractDate:
    def __init__(self):
        self.today = datetime.now()

    def extract_complete_date(self):
        date_extract = self.today.strftime("%Y-%m-%d_%H-%M-%S")
        
        return date_extract
    

    def extract_year_date(self):
        year = str(self.today.year)

        return year


    def extract_month_date(self):
        month = self.today.strftime('%m')

        return month


    def extract_day_date(self):
        day = self.today.strftime('%d')

        return day

    
    def extract_hour(self):
        hour = self.today.strftime("%H-%M-%S")

        return hour


class ExcelService():
    """
    Gera arquivo .xlsx e gerencia o rastreamento do arquivo gerado.\n
    O relatório será salvo em Histórico_UEX/ano/mes/dia/...
    """
    def __init__(self):
        self.date = ExctractDate()


    def create_path(self):
        path_root = pathlib.Path.cwd()

        path_records = path_root / 'Histórico_UEX' / self.date.extract_year_date() / self.date.extract_month_date() / self.date.extract_day_date()

        if os.path.isdir(path_records):
            pass
        else:
            path_records.mkdir(parents=True, exist_ok=True)

        return path_records


    def file_xlsx(self, title_parser, table_parser, name_file):
        self.file = Workbook()
        self.planilha = self.file.active
        
        
        if self.planilha is None:
            raise ValueError('Planilha não foi criada')
        
        # Cabeçalho - Identificação
        for i, linha in enumerate(title_parser, start=1):
            self.text = linha.get_text(strip=True)
            self.planilha.merge_cells(f'A{i}:H{i}')

            self.planilha[f'A{i}'] = self.text

        

        # Cabeçalho - Histórico
        header_columns_history = [
             'Data de envio',
             'Situação OPC UEx',
             'Situação PC EEx',
             'Situação OPC EEx',
             'Efeito Suspensivo UEx',
             'Efeito Suspensivo EEx',
             'Apta Pagamento',
             'Situação do Envio',
        ]
        
        for col, valor in enumerate(header_columns_history, start=1):
            self.planilha.cell(row=6, column=col, value=valor)


        # Dados - Histórico
        for row in table_parser:
            self.planilha.append(row)


        # Configura e adiciona estilo a tabela
        last_row = self.planilha.max_row
        table = Table(
            displayName='TabelaHistorico',
            ref=f'A6:H{last_row}',
            published=True
        )

        style = TableStyleInfo(
            name='TableStyleMedium8',
            showRowStripes=True
        )

        table.tableStyleInfo = style
        self.planilha.add_table(table)

        
        # Gera arquivo .xlsx
        path = self.create_path()

        hour = self.date.extract_hour()
        
        name_file_in_path = f'UEX - {name_file.text}_{hour}.xlsx'
        

        self.file.save(path / name_file_in_path)
